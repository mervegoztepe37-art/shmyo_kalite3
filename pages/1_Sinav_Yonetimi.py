import sqlite3, csv, io, re, hashlib
from docx import Document
import streamlit as st
import pandas as pd

DB='shmyo.db'
def conn():
    c=sqlite3.connect(DB)
    c.execute('PRAGMA foreign_keys=ON')
    return c

def init():
    with conn() as c:
        c.executescript('''
        CREATE TABLE IF NOT EXISTS programs(id INTEGER PRIMARY KEY,name TEXT UNIQUE NOT NULL);
        CREATE TABLE IF NOT EXISTS courses(id INTEGER PRIMARY KEY,program_id INTEGER NOT NULL REFERENCES programs(id),code TEXT NOT NULL,name TEXT NOT NULL,UNIQUE(program_id,code));
        CREATE TABLE IF NOT EXISTS outcomes(id INTEGER PRIMARY KEY,course_id INTEGER NOT NULL REFERENCES courses(id),code TEXT NOT NULL,description TEXT NOT NULL,UNIQUE(course_id,code));
        CREATE TABLE IF NOT EXISTS questions(id INTEGER PRIMARY KEY,course_id INTEGER NOT NULL REFERENCES courses(id),outcome_id INTEGER NOT NULL REFERENCES outcomes(id),body TEXT NOT NULL,a TEXT NOT NULL,b TEXT NOT NULL,c TEXT NOT NULL,d TEXT NOT NULL,e TEXT NOT NULL,correct TEXT NOT NULL CHECK(correct IN ('A','B','C','D','E')),bloom TEXT NOT NULL);
        CREATE TABLE IF NOT EXISTS exams(id INTEGER PRIMARY KEY,course_id INTEGER NOT NULL REFERENCES courses(id),name TEXT NOT NULL);
        CREATE TABLE IF NOT EXISTS exam_questions(exam_id INTEGER REFERENCES exams(id),question_id INTEGER REFERENCES questions(id),PRIMARY KEY(exam_id,question_id));
        CREATE TABLE IF NOT EXISTS responses(exam_id INTEGER REFERENCES exams(id),student_code TEXT NOT NULL,question_id INTEGER REFERENCES questions(id),answer TEXT CHECK(answer IN ('A','B','C','D','E') OR answer IS NULL),PRIMARY KEY(exam_id,student_code,question_id));
        ''')

def rows(sql,args=()):
    with conn() as c:
        return c.execute(sql,args).fetchall()

def execute(sql,args=()):
    with conn() as c:
        cur=c.execute(sql,args)
        return cur.lastrowid

def choices(table,extra=''):
    return rows(f'SELECT id,name FROM {table} {extra} ORDER BY name')


def normalize_question(s):
    return re.sub(r'\s+',' ',str(s).strip()).casefold()

def parse_docx_questions(upload):
    # Word paragraphs are read in document XML order; column/section ordering may vary.
    doc=Document(io.BytesIO(upload.getvalue()))
    lines=[]
    for para in doc.paragraphs:
        for line in para.text.splitlines():
            if line.strip():lines.append(line.strip())
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for line in cell.text.splitlines():
                    if line.strip():lines.append(line.strip())
    question_re=re.compile(r'^\s*(\d{1,3})\s*[).]\s*(.*)$')
    option_re=re.compile(r'^\s*([A-E])\s*[).]\s*(.*)$',re.I)
    result=[];current=None;active_option=None
    for line in lines:
        qm=question_re.match(line);om=option_re.match(line)
        if qm and 1<=int(qm.group(1))<=150:
            if current:result.append(current)
            current={'number':int(qm.group(1)),'body':qm.group(2).strip(),'options':{}};active_option=None
        elif om and current:
            active_option=om.group(1).upper()
            current['options'][active_option]=om.group(2).strip()
        elif current and active_option:
            current['options'][active_option]+=' '+line
        elif current:
            current['body']+=' '+line
    if current:result.append(current)
    for q in result:
        q['complete']=bool(q['body']) and all(q['options'].get(k,'').strip() for k in 'ABCDE')
    return result

st.set_page_config(page_title='SHMYO Ölçme ve Değerlendirme',layout='wide')
init()
st.title('SHMYO | Sınav ve Ölçme-Değerlendirme')
st.caption('Yerel prototip • SQLite veri tabanı • Gerçek öğrenci verilerini henüz yüklemeyin')
page=st.sidebar.radio('Modüller',['Programlar ve Dersler','Öğrenme Çıktıları','Soru Bankası','Sınav Hazırlama','Cevap Yükleme','Analiz ve Rapor'])

if page=='Programlar ve Dersler':
    st.subheader('Program ekle')
    with st.form('p'):
        name=st.text_input('Program adı')
        if st.form_submit_button('Programı kaydet') and name.strip():
            try: execute('INSERT INTO programs(name) VALUES(?)',(name.strip(),));st.success('Kaydedildi')
            except sqlite3.IntegrityError: st.error('Bu program zaten kayıtlı')
    ps=choices('programs')
    if ps:
        with st.form('course'):
            p=st.selectbox('Program',ps,format_func=lambda x:x[1]); code=st.text_input('Ders kodu'); n=st.text_input('Ders adı')
            if st.form_submit_button('Dersi kaydet') and code.strip() and n.strip():
                try: execute('INSERT INTO courses(program_id,code,name) VALUES(?,?,?)',(p[0],code.strip(),n.strip()));st.success('Ders eklendi')
                except sqlite3.IntegrityError: st.error('Bu ders kodu programda mevcut')
    st.dataframe(pd.DataFrame(rows('SELECT p.name,c.code,c.name FROM courses c JOIN programs p ON p.id=c.program_id'),columns=['Program','Ders kodu','Ders']),hide_index=True)
else:
    cs=rows('SELECT c.id,p.name||" / "||c.code||" - "||c.name FROM courses c JOIN programs p ON p.id=c.program_id ORDER BY p.name,c.name')
    if not cs: st.info('Önce bir program ve ders ekleyin.');st.stop()
    course=st.selectbox('Ders seçin',cs,format_func=lambda x:x[1]); cid=course[0]
    outs=rows('SELECT id,code||" - "||description FROM outcomes WHERE course_id=? ORDER BY code',(cid,))
    if page=='Öğrenme Çıktıları':
        with st.form('out'):
            code=st.text_input('DÖÇ kodu (ör. DÖÇ-1)'); desc=st.text_area('Öğrenme çıktısı')
            if st.form_submit_button('Kaydet') and code.strip() and desc.strip():
                try:execute('INSERT INTO outcomes(course_id,code,description) VALUES(?,?,?)',(cid,code.strip(),desc.strip()));st.success('Kaydedildi')
                except sqlite3.IntegrityError:st.error('Bu kod mevcut')
        st.dataframe(pd.DataFrame(rows('SELECT code,description FROM outcomes WHERE course_id=?',(cid,)),columns=['Kod','Açıklama']),hide_index=True)
    elif page=='Soru Bankası':
        if not outs:st.info('Önce ders öğrenme çıktısı ekleyin.');st.stop()
        with st.form('q'):
            body=st.text_area('Soru metni'); opts=[st.text_input(f'{k} şıkkı') for k in 'ABCDE']
            correct=st.selectbox('Doğru cevap',list('ABCDE')); out=st.selectbox('Öğrenme çıktısı',outs,format_func=lambda x:x[1]); bloom=st.selectbox('Bloom düzeyi',['Hatırlama','Anlama','Uygulama','Analiz','Değerlendirme','Yaratma'])
            if st.form_submit_button('Soruyu kaydet'):
                if body.strip() and all(v.strip() for v in opts):
                    execute('INSERT INTO questions(course_id,outcome_id,body,a,b,c,d,e,correct,bloom) VALUES(?,?,?,?,?,?,?,?,?,?)',(cid,out[0],body,*opts,correct,bloom));st.success('Soru kaydedildi')
                else:st.error('Soru ve tüm şıklar zorunlu')
        st.dataframe(pd.DataFrame(rows('SELECT q.id,q.body,o.code,q.bloom FROM questions q JOIN outcomes o ON o.id=q.outcome_id WHERE q.course_id=?',(cid,)),columns=['ID','Soru','DÖÇ','Bloom']),hide_index=True)
        st.divider()
        st.subheader('Word’den toplu soru ekle')
        st.caption('DOCX dosyasındaki soru ve A–E seçeneklerini ayırır. Her sorunun doğru cevabı ve DÖÇ bilgisi onaylanmadan kayıt yapılmaz. İki sütunlu veya birden fazla kitapçıklı belgelerde önizlemeyi mutlaka kontrol edin.')
        file=st.file_uploader('Word sınav kitapçığı (.docx)',type=['docx'],key='bank_word')
        if file is not None:
            fingerprint=hashlib.sha256(file.getvalue()).hexdigest()
            if st.session_state.get('word_file_fingerprint')!=fingerprint or st.session_state.get('word_course_id')!=cid:
                st.session_state.word_file_fingerprint=fingerprint
                st.session_state.word_course_id=cid
                st.session_state.word_parsed=parse_docx_questions(file)
            parsed=st.session_state.word_parsed
            st.info(f'{len(parsed)} soru adayı tespit edildi. Bu sayı birden fazla kitapçığı da içerebilir.')
            if parsed:
                df=pd.DataFrame([{'Aktar':q['complete'],'Soru no':q['number'],'Soru':q['body'],**{k:q['options'].get(k,'') for k in 'ABCDE'},'Doğru cevap':'','DÖÇ':outs[0][1].split(' - ')[0],'Bloom':'Hatırlama'} for q in parsed])
                edited=st.data_editor(df,hide_index=True,num_rows='fixed',use_container_width=True,key='word_editor',column_config={'Aktar':st.column_config.CheckboxColumn('Aktar'),'Doğru cevap':st.column_config.SelectboxColumn('Doğru cevap',options=['','A','B','C','D','E']),'DÖÇ':st.column_config.SelectboxColumn('DÖÇ',options=[x[1].split(' - ')[0] for x in outs]),'Bloom':st.column_config.SelectboxColumn('Bloom',options=['Hatırlama','Anlama','Uygulama','Analiz','Değerlendirme','Yaratma'])})
                st.caption('Aynı sorunun farklı kitapçıklardaki seçenek sırası değişebileceğinden doğru cevapları her satır için kontrol edin. Görseller bu sürümde aktarılmaz.')
                if st.button('Onaylanan soruları soru bankasına kaydet',type='primary'):
                    selected=edited[edited['Aktar']==True]
                    if selected.empty: st.warning('Aktarılacak soru seçilmedi.')
                    elif selected['Doğru cevap'].isin(list('ABCDE')).sum()!=len(selected):st.error('Her seçilen soru için doğru cevabı seçin.')
                    elif any(not str(r['Soru']).strip() or any(not str(r[k]).strip() for k in 'ABCDE') for _,r in selected.iterrows()):st.error('Seçilen soruların metin ve A–E seçenekleri eksiksiz olmalı.')
                    else:
                        outcome_map={x[1].split(' - ')[0]:x[0] for x in outs}
                        count=0;duplicates=0
                        with conn() as c:
                            existing={(normalize_question(r[0]),tuple(normalize_question(v) for v in r[1:])) for r in c.execute('SELECT body,a,b,c,d,e FROM questions WHERE course_id=?',(cid,)).fetchall()}
                            for _,r in selected.iterrows():
                                key=(normalize_question(r['Soru']),tuple(normalize_question(r[k]) for k in 'ABCDE'))
                                if key in existing:duplicates+=1;continue
                                c.execute('INSERT INTO questions(course_id,outcome_id,body,a,b,c,d,e,correct,bloom) VALUES(?,?,?,?,?,?,?,?,?,?)',(cid,outcome_map[r['DÖÇ']],r['Soru'],*[r[k] for k in 'ABCDE'],r['Doğru cevap'],r['Bloom']))
                                existing.add(key);count+=1
                        st.success(f'{count} soru kaydedildi. {duplicates} birebir tekrar atlandı. Sınav Hazırlama menüsünden seçebilirsiniz.')

    elif page=='Sınav Hazırlama':
        qs=rows('SELECT id,body FROM questions WHERE course_id=?',(cid,))
        if not qs:st.info('Önce soru ekleyin.');st.stop()
        with st.form('exam'):
            name=st.text_input('Sınav adı'); selected=st.multiselect('Sorular',qs,format_func=lambda x:f'{x[0]} | {x[1][:100]}')
            if st.form_submit_button('Sınavı oluştur') and name.strip() and selected:
                with conn() as c:
                    eid=c.execute('INSERT INTO exams(course_id,name) VALUES(?,?)',(cid,name.strip())).lastrowid
                    c.executemany('INSERT INTO exam_questions VALUES(?,?)',[(eid,q[0]) for q in selected])
                st.success(f'Sınav oluşturuldu. ID: {eid}')
    else:
        exams=rows('SELECT id,name FROM exams WHERE course_id=? ORDER BY id DESC',(cid,))
        if not exams:st.info('Önce sınav oluşturun.');st.stop()
        exam=st.selectbox('Sınav',exams,format_func=lambda x:f'{x[0]} - {x[1]}');eid=exam[0]
        questions=rows('SELECT q.id,q.correct,q.outcome_id,q.body FROM questions q JOIN exam_questions eq ON q.id=eq.question_id WHERE eq.exam_id=? ORDER BY q.id',(eid,))
        if page=='Cevap Yükleme':
            template=pd.DataFrame(columns=['student_code']+[str(q[0]) for q in questions])
            st.download_button('Boş cevap şablonunu indir',template.to_csv(index=False).encode('utf-8-sig'),'cevap_sablonu.csv','text/csv')
            st.caption('Her satır bir öğrenci; sütunlar soru ID; cevaplar A, B, C, D, E veya boş.')
            upload=st.file_uploader('Doldurulmuş CSV dosyası',type='csv')
            if upload:
                try:
                    df=pd.read_csv(upload,dtype=str,keep_default_na=False).fillna(''); expected=['student_code']+[str(q[0]) for q in questions]
                    if list(df.columns)!=expected:raise ValueError('CSV sütunları şablonla aynı olmalı')
                    if df.student_code.eq('').any() or df.student_code.duplicated().any():raise ValueError('Öğrenci kodları boş veya tekrar eden olamaz')
                    data=[]
                    for _,r in df.iterrows():
                        for q in questions:
                            a=r[str(q[0])].strip().upper()
                            if a not in ('','A','B','C','D','E'):raise ValueError(f'Geçersiz cevap: {a}')
                            data.append((eid,r['student_code'].strip(),q[0],a or None))
                    if st.button(f'{len(df)} öğrencinin cevaplarını kaydet'):
                        with conn() as c:
                            c.execute('DELETE FROM responses WHERE exam_id=?',(eid,))
                            c.executemany('INSERT INTO responses VALUES(?,?,?,?)',data)
                        st.success('Cevaplar kaydedildi')
                except Exception as e:st.error(str(e))
        else:
            rr=rows('SELECT student_code,question_id,answer FROM responses WHERE exam_id=?',(eid,))
            if not rr:st.info('Önce cevapları yükleyin.');st.stop()
            answer={(s,q):a for s,q,a in rr};students=sorted(set(s for s,_,_ in rr)); n=len(students)
            scores={s:sum(answer.get((s,q[0]))==q[1] for q in questions) for s in students}
            st.metric('Öğrenci sayısı',n);st.metric('Ortalama başarı (%)',round(100*sum(scores.values())/(n*len(questions)),2))
            st.dataframe(pd.DataFrame([{'Öğrenci':s,'Doğru':scores[s],'Başarı (%)':round(100*scores[s]/len(questions),2)} for s in students]),hide_index=True)
            analysis=[]
            for qid,correct,oid,body in questions:
                corrects=[int(answer.get((s,qid))==correct) for s in students]
                p=sum(corrects)/n
                # Corrected item-total correlation, undefined if either vector is constant.
                import numpy as np
                x=np.array(corrects,dtype=float);rest=np.array([scores[s]-corrects[i] for i,s in enumerate(students)],dtype=float)
                discr=float(np.corrcoef(x,rest)[0,1]) if n>=3 and x.std()>0 and rest.std()>0 else None
                analysis.append({'Soru ID':qid,'Soru':body[:70],'DÖÇ ID':oid,'Güçlük (p)':round(p,3),'Düzeltilmiş madde-toplam korelasyonu':round(discr,3) if discr is not None else None})
            adf=pd.DataFrame(analysis);st.subheader('Madde analizi');st.dataframe(adf,hide_index=True)
            st.download_button('Madde analizini CSV indir',adf.to_csv(index=False).encode('utf-8-sig'),'madde_analizi.csv','text/csv')
            outcome_stats=[]
            for oid in sorted(set(q[2] for q in questions)):
                relevant=[q for q in questions if q[2]==oid];hits=sum(answer.get((s,q[0]))==q[1] for s in students for q in relevant)
                outcome_stats.append({'DÖÇ ID':oid,'Soru sayısı':len(relevant),'Başarı (%)':round(100*hits/(n*len(relevant)),2)})
            st.subheader('Kazanım analizi');st.dataframe(pd.DataFrame(outcome_stats),hide_index=True)
            # KR-20 for dichotomous items, population variance convention.
            k=len(questions);var=pd.Series(list(scores.values()),dtype=float).var(ddof=0)
            if k>1 and n>=2 and var>0:
                kr20=k/(k-1)*(1-sum(a['Güçlük (p)']*(1-a['Güçlük (p)']) for a in analysis)/var)
                st.metric('KR-20 (yaklaşık)',round(kr20,3))
            else:st.caption('KR-20 için yeterli değişkenlik/veri yok.')
